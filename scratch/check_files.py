import sys
import os
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from docx import Document

# Check PPT
print("=== POWERPOINT ===")
prs = Presentation('Bao_Cao_Slide_Thuyet_Trinh_De_Tai_4.pptx')
print(f"Số slide: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    title = slide.shapes.title.text if slide.shapes.title else "(no title)"
    print(f"  Slide {i+1}: {title}")

print("\n=== WORD REPORT ===")
doc = Document('Bao_Cao_Do_An_Khai_Thac_Du_Lieu_De_Tai_4.docx')
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        print(f"  [{p.style.name}] {p.text}")
